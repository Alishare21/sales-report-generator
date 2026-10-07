"""Independent checks of the example's source calculations and export structure."""
from pathlib import Path
import hashlib
import json
import math
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]


def verify():
    folder=ROOT/'outputs'/'example_retail'
    evidence=json.loads((folder/'evidence.json').read_text(encoding='utf-8'))
    raw=ROOT/'data'/'raw'/'example_retail.csv'
    assert hashlib.sha256(raw.read_bytes()).hexdigest()==evidence['source_sha256']
    frame=pd.read_csv(raw).drop_duplicates()
    frame=frame.loc[(frame.quantity>=0)&(~(frame.unit_price<0))].copy()
    total=(frame.quantity*frame.unit_price).sum()
    means=frame.groupby(frame.promoted.str.lower()).quantity.mean()
    gap=(means['yes']-means['no'])/means['no']
    assert len(frame)==evidence['facts']['rows']
    assert math.isclose(total,evidence['facts']['total'],rel_tol=1e-12)
    f=next(x for x in evidence['findings'] if x['kind']=='bar')
    assert math.isclose(means['yes'],f['higher_mean'],rel_tol=1e-12)
    assert math.isclose(means['no'],f['lower_mean'],rel_tol=1e-12)
    for group in ['product_category','sales_region']:
        parts=(frame.quantity*frame.unit_price).groupby(frame[group].str.strip().str.casefold(),dropna=False).sum()
        assert math.isclose(parts.sum(),total,rel_tol=1e-12)
    book=load_workbook(folder/'findings_report.xlsx',data_only=False)
    assert book.sheetnames==['Summary','Findings','Charts','Recommendations','Data_Summary','Cleaned_Data','Notes']
    for sh in book:
        assert sh.freeze_panes=='A2'
        assert not sh.merged_cells.ranges
    assert book['Cleaned_Data'].auto_filter.ref
    assert len(book['Charts']._images)==len(evidence['findings'])
    formulas=sum(cell.data_type=='f' for row in book['Data_Summary'] for cell in row)
    assert formulas>0
    book.close()
    reader=PdfReader(folder/'findings_report.pdf')
    text='\n'.join(p.extract_text() for p in reader.pages)
    for finding in evidence['findings']:
        assert f"Figure {finding['figure']}" in text
    assert 'SYNTHETIC EXAMPLE' in text
    assert not any(s in text for s in ['[Your answer here]','TODO','{{'])
    first_appendix=next(i for i,p in enumerate(reader.pages) if 'Appendix:' in p.extract_text())
    assert first_appendix<=8
    checks={'source_rows_after_cleaning':len(frame),'estimated_revenue_inr':float(total),
            'promoted_mean_units':float(means['yes']),'nonpromoted_mean_units':float(means['no']),
            'relative_mean_gap':float(gap),'pdf_pages':len(reader.pages),'main_pages':first_appendix,
            'data_summary_formulas':formulas,'raw_unchanged':True}
    (folder/'independent_checks.json').write_text(json.dumps(checks,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(checks,indent=2))


if __name__=='__main__':verify()
