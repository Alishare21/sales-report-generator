"""Conservative tabular loading and profiling. Raw files are never opened for writing."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


class NeedsInput(ValueError):
    """A material interpretation needs a user decision before export."""


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def key(name):
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def _headers(values):
    if any(pd.isna(x) or not str(x).strip() for x in values):
        raise NeedsInput("Blank column headers: identify the table/header row first.")
    names = [str(x).strip().lstrip("\ufeff") for x in values]
    if len(set(names)) != len(names):
        raise NeedsInput("Duplicate column headers need distinct business meanings.")
    return names


def _json_frame(obj, config):
    route = config.get("json_records_key")
    if route:
        for part in route.split("."):
            obj = obj[part]
    elif isinstance(obj, dict):
        arrays = [v for v in obj.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        if len(arrays) == 1:
            obj = arrays[0]
        elif len(arrays) > 1:
            raise NeedsInput("Multiple JSON record arrays: set json_records_key.")
        elif all(isinstance(v, list) for v in obj.values()):
            return pd.DataFrame(obj)
        else:
            obj = [obj]
    if not isinstance(obj, list) or not all(isinstance(x, dict) for x in obj):
        raise NeedsInput("JSON must contain records (objects), not an unlabelled scalar array.")
    frame = pd.json_normalize(obj, sep=".")
    # Retain nested arrays as text; exploding would change the row grain.
    for col in frame:
        frame[col] = frame[col].map(lambda x: json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list)) else x)
    return frame


def load_table(path, config=None):
    config = config or {}
    path = Path(path)
    content = path.read_bytes()
    ext = path.suffix.lower()
    meta = {"source": str(path.resolve()), "sha256": digest(path), "format": ext,
            "warnings": [], "sheet_inventory": []}
    if not content:
        raise NeedsInput("The input file is empty.")
    if ext == ".xlsx":
        if not zipfile.is_zipfile(path):
            raise NeedsInput("XLSX extension does not match its ZIP-based content.")
        book = load_workbook(path, data_only=True, read_only=False)
        formula_book = load_workbook(path, data_only=False, read_only=False)
        candidates = []
        for sheet in book:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            occupied = sum(any(x is not None for x in r) for r in rows)
            meta["sheet_inventory"].append({"name": sheet.title, "nonempty_rows": occupied,
                                             "merged_ranges": [str(x) for x in sheet.merged_cells.ranges]})
            candidates.append((occupied, sheet.title, rows))
        chosen = config.get("sheet") or max(candidates, key=lambda x: x[0])[1]
        if chosen not in book.sheetnames:
            raise NeedsInput(f"Unknown sheet: {chosen}")
        rows = next(x[2] for x in candidates if x[1] == chosen)
        if not rows:
            raise NeedsInput("Selected sheet is empty.")
        if "header_row" in config:
            header = int(config["header_row"]) - 1
        else:
            possible = [(sum(x is not None for x in row), -i, i)
                        for i, row in enumerate(rows[:30])
                        if sum(x is not None for x in row) >= 2
                        and all(isinstance(x, str) for x in row if x is not None)
                        and len(set(x for x in row if x is not None)) == sum(x is not None for x in row)]
            if not possible:
                raise NeedsInput("Cannot identify Excel headers. Supply header_row (one-based).")
            header = max(possible)[2]
        active = [i for i, x in enumerate(rows[header]) if x is not None]
        if active != list(range(min(active), max(active) + 1)):
            raise NeedsInput("Noncontiguous Excel table; choose a clean table export.")
        names = _headers([rows[header][i] for i in active])
        for row in formula_book[chosen].iter_rows(min_row=header + 2):
            for cell in row:
                if cell.data_type == "f" and book[chosen][cell.coordinate].value is None:
                    raise NeedsInput(f"Formula {chosen}!{cell.coordinate} has no cached result. Recalculate and save in Excel.")
        frame = pd.DataFrame([[r[i] for i in active] for r in rows[header + 1:]], columns=names)
        meta.update(sheet=chosen, header_row=header + 1)
        book.close()
        formula_book.close()
    elif ext in (".csv", ".tsv", ".json"):
        if content.startswith(b"PK\x03\x04") or b"\x00" in content[:4096]:
            raise NeedsInput("Text extension has binary content (or unsupported UTF-16 encoding).")
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                text = content.decode(encoding).lstrip("\ufeff")
                break
            except UnicodeDecodeError:
                continue
        meta["encoding"] = encoding
        meta["physical_lines"] = len(text.splitlines())
        if ext == ".json":
            try:
                frame = _json_frame(json.loads(text), config)
            except json.JSONDecodeError as exc:
                raise NeedsInput("JSON extension does not match valid JSON content.") from exc
        else:
            if text.lstrip().startswith(("{", "[", "<!DOCTYPE", "<html")):
                raise NeedsInput("Delimited extension appears to contain JSON or HTML.")
            delimiters = [",", ";", "\t", "|"]
            try:
                guessed = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
                delimiters.remove(guessed)
                delimiters.insert(0, guessed)
            except csv.Error:
                pass
            candidates = []
            for delim in delimiters:
                parsed = list(csv.reader(io.StringIO(text), delimiter=delim, strict=True))
                parsed = [r for r in parsed if r]
                widths = Counter(len(r) for r in parsed)
                if parsed and len(widths) == 1:
                    candidates.append((len(parsed[0]), -delimiters.index(delim), delim, parsed))
            if not candidates:
                raise NeedsInput("Inconsistent field counts: check quoting/delimiter before analysis.")
            width, _, delim, parsed = max(candidates, key=lambda x: (x[0], x[1]))
            if width == 1 and any(d in parsed[0][0] for d in delimiters):
                raise NeedsInput("One giant column contains separators; delimiter unresolved.")
            names = _headers(parsed[0])
            frame = pd.DataFrame(parsed[1:], columns=names)
            frame = frame.replace(r"^\s*$", pd.NA, regex=True)
            meta.update(delimiter=repr(delim), logical_records=len(parsed) - 1)
            if len(frame) != meta["logical_records"]:
                raise NeedsInput("Parsed record count failed sanity check.")
    else:
        raise NeedsInput(f"Unsupported extension {ext}; use CSV, TSV, XLSX or JSON.")
    frame.columns = _headers(frame.columns)
    if frame.empty or len(frame.columns) == 0:
        raise NeedsInput("No data records found.")
    meta.update(shape=list(frame.shape), columns=list(frame.columns),
                dtypes={c: str(frame[c].dtype) for c in frame},
                first_five=frame.head(5).astype(object).where(frame.head(5).notna(), None).to_dict("records"))
    return frame, meta


def classify(frame, config=None):
    config = config or {}
    result, date_formats = {}, dict(config.get("date_formats", {}))
    numeric_roles = set(v for k, v in config.get("roles", {}).items() if k in {"quantity", "price", "revenue", "cost", "profit"})
    for col in frame:
        s = frame[col].dropna()
        n, unique = len(s), s.nunique()
        name = key(col)
        if col in config.get("classes", {}):
            kind = config["classes"][col]
        elif unique <= 1:
            kind = "Constant"
        elif name == "id" or name.endswith("_id") or name in {"sku", "barcode", "postcode", "zip_code"}:
            kind = "ID"
        elif pd.api.types.is_datetime64_any_dtype(s):
            kind = "Date/time"
            date_formats.setdefault(col, "ISO8601")
        elif col in date_formats:
            kind = "Date/time"
        elif name in {'promo','promoted','promotion','campaign_flag','ab_group'} and unique == 2:
            kind = "Category"
        elif s.astype(str).str.match(r"^\d{4}-\d{2}-\d{2}(?:[ T].*)?$").mean() >= 0.8:
            kind = "Date/time"
            date_formats.setdefault(col, "ISO8601")
        elif name.endswith("date") or name in {"date", "timestamp", "datetime", "time"}:
            kind = "Date/time"
            # Refuse ambiguous day/month order, even when a library would guess.
        elif col in numeric_roles or pd.to_numeric(s, errors="coerce").notna().mean() >= 0.9:
            kind = "Numeric measure"
        elif s.astype(str).str.match(r"^(?:https?://|www\.)").mean() > 0.5 or any(x in name.split("_") for x in ("name", "description", "url", "email", "comment", "notes")):
            kind = "Free text / URL"
        elif unique <= min(30, max(2, n * 0.2)):
            kind = "Category"
        elif unique / max(1, n) >= 0.98 and s.astype(str).str.fullmatch(r"[A-Za-z0-9_-]{1,40}").all():
            kind = "ID"
        else:
            kind = "Free text / URL"
        if kind not in {"ID", "Date/time", "Numeric measure", "Category", "Free text / URL", "Constant"}:
            raise NeedsInput(f"Invalid classification for {col}: {kind}")
        result[col] = kind
    return result, date_formats


def profile(frame, classes, formats, as_of):
    columns = []
    for col in frame:
        s = frame[col]
        values = s.dropna()
        normalized = values.astype(str).str.strip().str.casefold()
        item = {"column": col, "dtype": str(s.dtype), "class": classes[col],
                "null_count": int(s.isna().sum()), "null_pct": float(s.isna().mean() * 100),
                "unique_count": int(s.nunique()), "duplicate_id_rows": int(values.duplicated().sum()) if classes[col] == "ID" else None,
                "case_whitespace_variants": int(values.astype(str).nunique() - normalized.nunique()) if classes[col] == "Category" else 0,
                "whitespace_cells": int((values.astype(str) != values.astype(str).str.strip()).sum()) if classes[col] == "Category" else 0}
        if classes[col] == "Numeric measure":
            numeric = pd.to_numeric(s, errors="coerce")
            finite = numeric.replace([np.inf, -np.inf], np.nan)
            item.update(min=finite.min(), max=finite.max(), negative_count=int((numeric < 0).sum()),
                        parse_failures=int((s.notna() & numeric.isna()).sum()), nonfinite_count=int(np.isinf(numeric.astype(float)).sum()))
        elif classes[col] == "Date/time" and col in formats:
            dates = pd.to_datetime(s, format=formats[col], errors="coerce", utc=True)
            item.update(min=str(dates.min()), max=str(dates.max()), future_count=int((dates.dt.date > as_of).sum()),
                        parse_failures=int((s.notna() & dates.isna()).sum()))
        else:
            item["top_values"] = {str(k): int(v) for k, v in values.value_counts().head(5).items()}
        columns.append(item)
    return {"columns": columns, "exact_duplicate_rows": int(frame.duplicated().sum())}
