import json
from pathlib import Path
import sys
import tempfile
import unittest
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from build_report import build_excel


class ExportTests(unittest.TestCase):
    def test_literal_formula_text_and_wrapped_rows(self):
        data=pd.DataFrame({'record_id':['=1+1','x'],'quantity':[1.,2.],'notes':['A fairly long note that must stay fully visible in the source table.','short']})
        report={'dataset':'test.csv','generated':'2026-10-07','synthetic':False,
            'facts':{'rows':2,'period':'not available','total':3.,'valid_metric_rows':2,'mean':1.5},
            'capabilities':{'currency':None,'metric':'quantity','metric_label':'quantity','unit':'units','roles':{'quantity':'quantity'},'assumptions':[],'skipped':[]},
            'definitions':{'Primary total':'Sum of quantity','Observed metric records':'Nonmissing quantity count'},
            'findings':[],'recommendations':[],'classes':{'notes':'Free text / URL'},'groups':{},'source':'test.csv',
            'source_sha256':'test','load':{'shape':[2,3]},'limitations':[],'cleaning_log':[],'profile':{'columns':[]}}
        with tempfile.TemporaryDirectory() as folder:
            target=Path(folder)/'test.xlsx'
            build_excel(report,data,target)
            book=load_workbook(target,data_only=False)
            self.assertEqual(book['Cleaned_Data']['A2'].value,'=1+1')
            self.assertEqual(book['Cleaned_Data']['A2'].data_type,'s')
            self.assertGreater(book['Cleaned_Data'].row_dimensions[2].height,22)
            self.assertEqual(book['Data_Summary']['B3'].data_type,'f')
            self.assertIn('[$-409]',book['Summary']['B9'].number_format)
            book.close()
            cached=load_workbook(target,data_only=True)
            self.assertEqual(cached['Data_Summary']['B3'].value,3.)
            cached.close()


if __name__=='__main__':unittest.main()
