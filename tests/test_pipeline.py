import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from load import load_table, classify, profile, digest, NeedsInput
from analysis import discover, clean, capabilities, analyze


class LoaderTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.folder=Path(self.temp.name)
    def tearDown(self):
        self.temp.cleanup()
    def write(self,name,text,encoding='utf-8'):
        path=self.folder/name;path.write_text(text,encoding=encoding);return path
    def test_delimiters_bom_and_leading_zero(self):
        for delimiter in [',',';','\t','|']:
            path=self.write('input.csv','\ufeff'+delimiter.join(['item_id','quantity'])+'\n'+delimiter.join(['001','5'])+'\n')
            frame,meta=load_table(path)
            self.assertEqual(frame.shape,(1,2));self.assertEqual(frame.iloc[0,0],'001')
    def test_quoted_newlines(self):
        frame,meta=load_table(self.write('input.csv','id,note\n1,"hello\nworld"\n2,ok\n'))
        self.assertEqual(len(frame),2);self.assertEqual(meta['physical_lines'],4)
    def test_latin1(self):
        frame,meta=load_table(self.write('input.csv','category;quantity\ncafé;2\n',encoding='latin-1'))
        self.assertEqual(frame.iloc[0,0],'café');self.assertEqual(meta['encoding'],'latin-1')
    def test_extension_mismatch(self):
        with self.assertRaises(NeedsInput):load_table(self.write('input.csv','[{"a":1}]'))
    def test_duplicate_headers(self):
        with self.assertRaises(NeedsInput):load_table(self.write('input.csv','a,a\n1,2\n'))
    def test_nested_json(self):
        path=self.write('input.json',json.dumps({'records':[{'id':'01','stats':{'score':3},'tags':['a','b']}]}))
        frame,_=load_table(path)
        self.assertIn('stats.score',frame);self.assertEqual(len(frame),1);self.assertEqual(frame.loc[0,'tags'],'["a", "b"]')
    def test_ambiguous_json(self):
        with self.assertRaises(NeedsInput):load_table(self.write('input.json','{"a":[{"x":1}],"b":[{"x":2}]}'))
    def test_excel_preamble_and_largest_sheet(self):
        path=self.folder/'input.xlsx';book=Workbook();book.active.title='Instructions';book.active.append(['Read me'])
        sh=book.create_sheet('Data');sh.merge_cells('A1:C1');sh['A1']='Sales extract';sh.append([])
        sh.append(['record_id','category','quantity']);sh.append(['r1','A',2]);sh.append(['r2','B',3]);sh.append(['Total',None,5]);book.save(path)
        frame,meta=load_table(path)
        self.assertEqual(meta['sheet'],'Data');self.assertEqual(meta['header_row'],3);self.assertEqual(len(frame),3)
        classes,fmts=classify(frame);roles=discover(frame,classes,{})
        data,log,_,blocks=clean(frame,classes,fmts,roles,{'remove_total_rows':True},date(2026,10,7))
        self.assertEqual(len(data),2);self.assertFalse(blocks)
    def test_excel_missing_formula_cache_stops(self):
        path=self.folder/'input.xlsx';book=Workbook();sh=book.active;sh.append(['id','quantity']);sh.append(['a','=1+1']);book.save(path)
        with self.assertRaises(NeedsInput):load_table(path)


class AnalysisTests(unittest.TestCase):
    def analyze(self,frame,config=None):
        config=config or {};classes,formats=classify(frame,config);roles=discover(frame,classes,config)
        data,log,limitations,blocks=clean(frame,classes,formats,roles,config,date(2026,10,7))
        cap=capabilities(data,classes,roles,config)
        return data,log,limitations,blocks,cap,classes
    def test_numeric_unique_values_are_not_ids(self):
        frame=pd.DataFrame({'score':np.linspace(.1,5,100),'record_id':[f'x{i}' for i in range(100)]})
        cls,_=classify(frame);self.assertEqual(cls['score'],'Numeric measure');self.assertEqual(cls['record_id'],'ID')
    def test_exact_duplicates_only(self):
        frame=pd.DataFrame({'item_id':['a','a','a'],'quantity':['1','2','2'],'url':['https://example.invalid']*3})
        data,log,_,_,_,_=self.analyze(frame)
        self.assertEqual(len(data),2);self.assertEqual(data['quantity'].sum(),3)
        self.assertEqual(log[0]['rows_before']-log[-1]['rows_after'],1)
    def test_negative_values_require_interpretation(self):
        frame=pd.DataFrame({'quantity':[-1,2,3,4]})
        self.assertTrue(self.analyze(frame)[3])
        result=self.analyze(frame,{'nonnegative':['quantity']});self.assertEqual(len(result[0]),3);self.assertFalse(result[3])
    def test_ambiguous_dates_block(self):
        frame=pd.DataFrame({'date':['01/02/2026','02/03/2026'],'quantity':[1,2]})
        self.assertTrue(self.analyze(frame)[3])
        self.assertFalse(self.analyze(frame,{'date_formats':{'date':'%d/%m/%Y'}})[3])
    def test_future_dates_block(self):
        frame=pd.DataFrame({'date':['2027-01-01','2026-01-01'],'quantity':[1,2]})
        self.assertTrue(self.analyze(frame)[3])
    def test_no_fabricated_dimensions_or_amounts(self):
        frame=pd.DataFrame({'department':['a']*40+['b']*40,'rating':list(range(80))})
        data,_,_,_,cap,cls=self.analyze(frame)
        self.assertIsNone(cap['metric']);self.assertNotIn('date',cap['roles']);self.assertNotIn('profit',cap['roles'])
        facts,*_=analyze(data,cap,cls,{})
        self.assertEqual(facts['total'],80);self.assertEqual(facts['period'],'not available')
    def test_mixed_currencies_block_amount_aggregation(self):
        frame=pd.DataFrame({'revenue':[10,20,30,40],'currency':['USD','EUR','USD','EUR']})
        self.assertIsNotNone(self.analyze(frame)[4]['currency_problem'])
    def test_low_n_no_group_claim(self):
        frame=pd.DataFrame({'category':['a']*29+['b']*40,'quantity':list(range(69))})
        data,_,_,_,cap,cls=self.analyze(frame)
        _,findings,_,_,_=analyze(data,cap,cls,{})
        self.assertFalse(any(f['kind']=='barh' for f in findings))
    def test_category_normalization_and_missing_reconciliation(self):
        frame=pd.DataFrame({'record_id':[f'x{i}' for i in range(120)],'category':[' A ','a','B',None]*30,'quantity':[1,2,3,4]*30})
        data,_,_,_,cap,cls=self.analyze(frame)
        facts,_,_,groups,_=analyze(data,cap,cls,{})
        self.assertEqual(sum(x['sum'] for x in groups['category']),facts['total'])
        self.assertEqual(data['category'].nunique(),2)
    def test_conflicting_row_ids_stop(self):
        frame=pd.DataFrame({'record_id':['a','a','b'],'quantity':[1,2,3]})
        self.assertTrue(self.analyze(frame,{'row_key':'record_id'})[3])
    def test_missing_input_is_not_zero_revenue(self):
        frame=pd.DataFrame({'record_id':['a','b','c'],'quantity':[1,2,3],'unit_price':[10,None,30]})
        data,_,_,_,cap,_=self.analyze(frame)
        self.assertEqual(data['revenue_est'].sum(),100);self.assertTrue(pd.isna(data.loc[1,'revenue_est']))
    def test_constant_price_can_derive_amount_but_is_not_charted(self):
        frame=pd.DataFrame({'quantity':[1,2,3],'unit_price':[10,10,10]})
        data,_,_,_,cap,cls=self.analyze(frame)
        self.assertEqual(cls['unit_price'],'Constant');self.assertEqual(data['revenue_est'].sum(),60)
        self.assertNotIn('unit_price',cap['numeric'])
    def test_record_primary_prevents_unverified_revenue(self):
        frame=pd.DataFrame({'quantity':[1,2,3],'price':[5,6,7]})
        data,_,_,_,cap,cls=self.analyze(frame,{'primary_measure':'records'})
        self.assertIsNone(cap['metric']);self.assertNotIn('revenue_est',data)
        facts,*_=analyze(data,cap,cls,{'primary_measure':'records'})
        self.assertEqual(facts['total'],3)
    def test_multigroup_flat_comparison_and_selected_categories(self):
        frame=pd.DataFrame({'record_id':[str(i) for i in range(120)],'position':['a']*40+['b']*40+['c']*40,
                            'quantity':[100]*40+[101]*40+[102]*40,'type':['x','y']*60})
        cfg={'category_columns':['type'],'comparison_columns':['position'],'segment_columns':[],
             'finding_order':['position','type'],'allow_impact_projection':False}
        data,_,_,_,cap,cls=self.analyze(frame,cfg)
        _,findings,recs,_,_=analyze(data,cap,cls,cfg)
        self.assertEqual(findings[0]['group'],'position');self.assertTrue(findings[0]['small'])
        self.assertEqual(len(findings[0]['group_stats']),3)
        self.assertNotIn('Illustrative effect range',recs[0]['impact'])
        self.assertNotIn('100 records per arm',recs[0]['action'])
        self.assertIn('minimum detectable effect',recs[0]['impact'])
    def test_invalid_selected_category_is_rejected(self):
        with self.assertRaises(NeedsInput):self.analyze(pd.DataFrame({'quantity':[1,2,3]}),{'category_columns':['missing']})
    def test_within_segment_reversal_is_reported(self):
        rows=[]
        for segment,flag,n,value in [('a','yes',120,10),('a','no',30,12),('b','yes',30,2),('b','no',120,3)]:
            rows += [{'record_id':str(len(rows)+i),'segment':segment,'promoted':flag,'quantity':value} for i in range(n)]
        data,_,_,_,cap,cls=self.analyze(pd.DataFrame(rows))
        _,findings,_,_,_=analyze(data,cap,cls,{})
        finding=next(f for f in findings if f['kind']=='bar')
        self.assertEqual(len(finding['subsegments']),2)
        self.assertTrue(all(s['difference']<0 for s in finding['subsegments']))


if __name__=='__main__':unittest.main()
